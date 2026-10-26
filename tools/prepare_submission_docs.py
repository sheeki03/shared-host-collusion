"""Package a filled report as a readable manuscript, supplement and complete table archive.

Uses an Apart template DOCX as Pandoc's reference and retains all report prose. Long or wide generated tables move
to explicitly referenced Markdown/CSV files rather than being compressed into unreadable Word cells. No raw logs,
attempt-level diagnostics, coder keys or human-check rows are copied. This is formatting, not analysis.

Run with a document-capable Python providing python-docx and lxml, plus the Pandoc executable.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
import shutil
import subprocess
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def preserve_template_package(path: Path, reference: Path, pandoc: str = "pandoc") -> None:
    """Repair Pandoc's template omissions: embedded font MIME types and missing referenced table/paragraph styles.

    The Google template has no Pandoc Compact/Table styles. Without their definitions LibreOffice can flatten table
    text outside its cells. Existing reference styles are retained; only absent converter styles are added.
    """
    from lxml import etree
    ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    with ZipFile(reference) as z:
        ref = etree.fromstring(z.read("[Content_Types].xml"))
        defaults = {e.get("Extension"): e.get("ContentType") for e in ref if e.tag == f"{{{ns}}}Default"}
    with ZipFile(path) as z:
        entries = [(info, z.read(info.filename)) for info in z.infolist()]
    content = etree.fromstring(dict((info.filename, b) for info, b in entries)["[Content_Types].xml"])
    have = {e.get("Extension") for e in content if e.tag == f"{{{ns}}}Default"}
    needed = {Path(info.filename).suffix.lstrip(".") for info, _ in entries if info.filename.startswith("word/fonts/")}
    for extension in needed - have:
        if extension not in defaults:
            raise ValueError(f"No template content type for embedded font extension {extension}")
        etree.SubElement(content, f"{{{ns}}}Default", Extension=extension, ContentType=defaults[extension])
    default_doc = subprocess.run([pandoc, "--print-default-data-file", "reference.docx"], check=True,
                                 stdout=subprocess.PIPE).stdout
    with ZipFile(io.BytesIO(default_doc)) as z:
        default_styles = etree.fromstring(z.read("word/styles.xml"))
    parts = dict((info.filename, b) for info, b in entries)
    styles = etree.fromstring(parts["word/styles.xml"])
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ids = {e.get(f"{{{w}}}styleId") for e in styles if e.tag == f"{{{w}}}style"}
    for style in default_styles:
        if style.tag == f"{{{w}}}style" and style.get(f"{{{w}}}styleId") not in ids:
            styles.append(style)
    replacements = {"[Content_Types].xml": etree.tostring(content, xml_declaration=True, encoding="UTF-8"),
                    "word/styles.xml": etree.tostring(styles, xml_declaration=True, encoding="UTF-8")}
    temp = path.with_suffix(".types.tmp")
    with ZipFile(temp, "w", ZIP_DEFLATED) as z:
        for info, data in entries:
            z.writestr(info, replacements.get(info.filename, data))
    temp.replace(path)


def table_rows(text: str) -> list[list[str]]:
    rows = []
    for line in text.splitlines():
        if not line.startswith("|"):
            continue
        cells = [s.strip().replace(r"\|", "|") for s in re.split(r"(?<!\\)\|", line.strip().strip("|"))]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def external_tables(text: str, source: str, tables: Path, out: Path, tables_url: str | None = None) -> tuple[str, list[dict]]:
    """Preserve only explicitly cited generated tables; select print placement without changing values."""
    records = []
    for name in dict.fromkeys(re.findall(r"\{\{TABLE:([^}]+)\}\}", source)):
        p = tables / f"{name}.md"
        if not p.exists():
            raise ValueError(f"Missing cited table: {p}")
        body = p.read_text().strip()
        if body not in text:
            raise ValueError(f"Filled report does not contain current table {name}; regenerate it first")
        rows = table_rows(body)
        ncol = max(map(len, rows), default=0)
        longest = max((len(c) for row in rows for c in row), default=0)
        # A4's matrix has long environment descriptions; a transposed prose summary remains in Appendix H.
        moved = len(rows) > 35 or ncol > 8 or longest > 95
        target = out / "tables" / p.name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(p, target)
        csv_path = target.with_suffix(".csv")
        with csv_path.open("w", newline="") as f:
            csv.writer(f).writerows(rows)
        if moved:
            if tables_url:
                base = tables_url.rstrip("/")
                note = (f"The complete table is available as [{name} (Markdown)]({base}/{name}.md) and "
                        f"[{name} (CSV)]({base}/{name}.csv), and in the accompanying table archive. "
                        "All values are retained there at their generated precision.")
            else:
                note = (f"The complete table is supplied as **tables/{name}.md** and **tables/{name}.csv** in the "
                        "accompanying table archive. All values are retained there at their generated precision.")
            text = text.replace(body, note)
        records.append({"name": name, "rows_including_headers": len(rows), "columns": ncol,
                        "placement": "external table archive" if moved else "supplement and table archive",
                        "markdown_sha256": digest(target), "csv_sha256": digest(csv_path)})
    return text, records


def remove_theme_font(run) -> None:
    from docx.oxml.ns import qn
    run.font.name = "Old Standard TT"
    rf = run._r.get_or_add_rPr().find(qn("w:rFonts"))
    if rf is not None:
        for k in list(rf.attrib):
            if "theme" in k.lower():
                del rf.attrib[k]


def format_docx(path: Path, *, supplement: bool = False, combined: bool = False) -> None:
    """Use the official template's page and type system, with readable research tables and figures."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from docx.text.paragraph import Paragraph
    doc = Document(path)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
        sec.header_distance = sec.footer_distance = Inches(.5)
    for sty in doc.styles:
        if sty.type == 1:
            sty.font.name = "Old Standard TT"
            sty.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        try:
            sty = doc.styles[name]
        except KeyError:
            continue
        sty.font.name = "Old Standard TT"
        sty.font.size = Pt(11)
        sty.paragraph_format.line_spacing = 1.0
        sty.paragraph_format.space_before = Pt(0)
        sty.paragraph_format.space_after = Pt(2)
        sty.paragraph_format.widow_control = True
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 14), ("Heading 3", 12)):
        try:
            sty = doc.styles[name]
        except KeyError:
            continue
        sty.font.size = Pt(size)
        sty.font.bold = True
        sty.font.color.rgb = RGBColor(0, 0, 0)
        sty.paragraph_format.space_before = Pt(8 if name != "Title" else 0)
        sty.paragraph_format.space_after = Pt(3)
        sty.paragraph_format.line_spacing = 1.0
        sty.paragraph_format.keep_with_next = True
    first = True
    in_references = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if first and text:
            p.style = doc.styles["Title"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first = False
        elif p.style.name in ("Heading 1", "Heading 2"):
            p.style = doc.styles["Heading 2"]
        if text == "References":
            in_references = True
            p.paragraph_format.page_break_before = False
        if combined and text.startswith("Appendix A."):
            p.paragraph_format.page_break_before = True
        if text in ("Code and Data", "Supplementary material", "LLM Usage Statement"):
            in_references = False
        if p.style.name == "Source Code":
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(8)
        else:
            for run in p.runs:
                if run.style is None or run.style.name not in ("Verbatim Char", "Source Code"):
                    remove_theme_font(run)
                if in_references and p.style.name not in ("Heading 1", "Heading 2"):
                    run.font.size = Pt(10)
        # Keep a caption with the preceding figure when possible, without stretching images.
        if text.startswith("Figure ") or text.startswith("Table "):
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.keep_with_next = False
            for run in p.runs:
                run.font.size = Pt(9)
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
    # Figures fit the printable width. A taller appendix figure must retain readable labels; the two main figures
    # are wide enough that this height cap does not alter their page budget.
    for shape in doc.inline_shapes:
        factor = min(1, Inches(6.5)/shape.width, Inches(4.0)/shape.height)
        shape.width, shape.height = int(shape.width*factor), int(shape.height*factor)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        count = max(len(row.cells) for row in table.rows)
        if not count:
            continue
        widths = [6.5/count] * count
        if count >= 4:
            widths = [1.65] + [4.85/(count-1)]*(count-1)
        headings = [c.text.strip() for c in table.rows[0].cells]
        if headings[:2] == ["Arm", "Scratch"] and count == 5:
            widths = [1.65, .60, 1.80, 1.60, .85]
        elif count == 5 and "nanoGPT model ID" in headings:
            widths = [1.25, 1.40, 2.00, .95, .90]
        elif headings == ["ID", "Hypothesis", "Test"]:
            widths = [.45, 3.00, 3.05]
        for i, column in enumerate(table.columns):
            column.width = Inches(widths[min(i, count-1)])
        for rownum, row in enumerate(table.rows):
            if row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is None:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if rownum == 0:
                header = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(header)
            for i, cell in enumerate(row.cells):
                cell.width = Inches(widths[min(i, count-1)])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    el = OxmlElement(f"w:{side}")
                    for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        el.set(qn(f"w:{key}"), value)
                    borders.append(el)
                pr.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), "60"); el.set(qn("w:type"), "dxa")
                    margins.append(el)
                pr.append(margins)
                if rownum == 0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E8EDF1"); pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    if i > 0 and re.fullmatch(r"[-−+–\d\s.eE%,/()\[\]]+", p.text):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        remove_theme_font(run)
                        run.font.size = Pt(8 if count >= 6 else 9)
                        run.font.color.rgb = RGBColor(0, 0, 0)
        next_element = table._tbl.getnext()
        if next_element is not None and next_element.tag == qn("w:p"):
            Paragraph(next_element, doc._body).paragraph_format.space_before = Pt(5)
    doc.core_properties.author = "Shivom Sharma"
    doc.core_properties.last_modified_by = "Shivom Sharma"
    doc.save(path)


def apply_template_title_block(path: Path, reference: Path) -> None:
    """Replace Pandoc's plain front matter with the official template's own title block, filled in.

    The template's first-page block (ruled title with the sprint footnote, centred author grid, "With Apart Research",
    centred Abstract heading and indented abstract) is copied from the reference DOCX. The author takes the centre
    cell, the unused author cells and second author row are removed, and the abstract text replaces the italic
    guidance. The template's information box and guidance text are not copied. The text is unchanged.
    """
    import copy
    from lxml import etree
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    PR = "http://schemas.openxmlformats.org/package/2006/relationships"
    W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
    q = lambda tag: f"{{{W}}}{tag}"
    text = lambda el: "".join(t.text or "" for t in el.iter(q("t")))
    with ZipFile(reference) as z:
        tbody = etree.fromstring(z.read("word/document.xml")).find(q("body"))
        tnotes = etree.fromstring(z.read("word/footnotes.xml"))
        trels = etree.fromstring(z.read("word/_rels/footnotes.xml.rels"))
    with ZipFile(path) as z:
        entries = [(info, z.read(info.filename)) for info in z.infolist()]
    parts = {info.filename: data for info, data in entries}
    doc = etree.fromstring(parts["word/document.xml"])
    body = doc.find(q("body"))
    # Pandoc front matter: title, author line, "With Apart Research", sprint line, Abstract heading, abstract.
    front = []
    for el in body:
        if el.tag == q("p") and re.match(r"1\.\s*Introduction$", text(el).strip()):
            break
        front.append(el)
    else:
        raise ValueError("No '1. Introduction' heading after the front matter")
    paras = [p for p in front if p.tag == q("p") and text(p).strip()]
    title = text(paras[0]).strip()
    author = paras[1]
    name = "".join(text(r) for r in author.iter(q("r")) if r.find(f"{q('rPr')}/{q('b')}") is not None).strip()
    affiliation = text(author).replace(name, "", 1).strip()
    heading = next(i for i, p in enumerate(paras) if text(p).strip() == "Abstract")
    abstract = [text(p).strip() for p in paras[heading + 1:]]
    if not (title and name and affiliation and abstract) or [text(p).strip() for p in paras[2:3]] != ["With Apart Research"]:
        raise ValueError("Unexpected front matter; expected title, author, With Apart Research and Abstract")
    block = []
    for el in tbody:
        if el.tag == q("tbl") and text(el).startswith("How to use this template"):
            break
        # The empty spacer paragraphs around the block are omitted to save page length; the block itself is kept.
        if el.tag == q("tbl"):
            block.append(copy.deepcopy(el))
    else:
        raise ValueError("Template information box not found")

    def cell_texts(cell, values):
        # The template writes "Author name N", a line break and "Affiliation" as separate text nodes of one run.
        ts = [t for t in cell.iter(q("t")) if (t.text or "").strip()]
        if len(ts) != len(values):
            raise ValueError("Unexpected template author cell")
        for t, value in zip(ts, values):
            t.text = value

    def para_with(prefix):
        return next(p for el in block for p in el.iter(q("p")) if text(p).strip().startswith(prefix))

    title_para = para_with("PROJECT TITLE")
    next(t for t in title_para.iter(q("t")) if t.text and "PROJECT TITLE" in t.text).text = title
    # A long title at the template's 20 pt leaves one word on a third line; 18 pt keeps it to two lines.
    for size in title_para.iter(q("sz"), q("szCs")):
        if len(title) > 80:
            size.set(q("val"), "36")
    cells = {n: next(a for a in para_with(f"Author name {n}").iterancestors(q("tc"))) for n in range(1, 7)}
    row2 = next(cells[4].iterancestors(q("tr")))
    if any(c in row2.iter(q("tc")) for c in (cells[1], cells[2], cells[3])):
        raise ValueError("Template author rows are not separable")
    row2.getparent().remove(row2)
    cell_texts(cells[2], [name, affiliation])
    # One author: the centre cell becomes the author grid's only column, at the grid's full width, so a narrow
    # template column cannot wrap the name.
    grid = next(cells[2].iterancestors(q("tbl")))
    width = next(grid.iterancestors(q("tbl"))).find(f"{q('tblPr')}/{q('tblW')}").get(q("w"))
    for n in (1, 3):
        cells[n].getparent().remove(cells[n])
    for cols in grid.iter(q("tblGrid")):
        for col in cols.findall(q("gridCol")):
            cols.remove(col)
        etree.SubElement(cols, q("gridCol")).set(q("w"), width.split(".")[0])
    grid.find(f"{q('tblPr')}/{q('tblW')}").set(q("w"), width)
    tcpr = cells[2].find(q("tcPr"))
    for el in tcpr.findall(q("tcW")) + tcpr.findall(q("gridSpan")):
        tcpr.remove(el)
    tcw = etree.Element(q("tcW"))
    tcw.set(q("w"), width.split(".")[0])
    tcw.set(q("type"), "dxa")
    tcpr.insert(0, tcw)
    guide = para_with("Summarize your project")
    runs = guide.findall(q("r"))
    rpr = copy.deepcopy(runs[0].find(q("rPr")))
    for tag in ("i", "iCs", "u"):
        for el in rpr.findall(q(tag)):
            rpr.remove(el)
    previous = guide
    for k, para_text in enumerate(abstract):
        p = guide if k == 0 else copy.deepcopy(guide)
        for r in p.findall(q("r")) + p.findall(q("hyperlink")):
            p.remove(r)
        r = etree.SubElement(p, q("r"))
        r.append(copy.deepcopy(rpr))
        t = etree.SubElement(r, q("t"))
        t.text = para_text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if k:
            previous.addnext(p)
            previous = p
    for el in block:
        for b in list(el.iter(q("bookmarkStart"), q("bookmarkEnd"))):
            b.getparent().remove(b)
        for node in el.iter():
            for attr in (f"{{{W14}}}paraId", f"{{{W14}}}textId"):
                node.attrib.pop(attr, None)
    # The sprint footnote keeps the template's text and link under ids that do not collide with Pandoc's.
    notes = etree.fromstring(parts["word/footnotes.xml"])
    rels = etree.fromstring(parts.get("word/_rels/footnotes.xml.rels")
                            or f'<Relationships xmlns="{PR}"/>'.encode())
    new_id = str(max(int(n.get(q("id"))) for n in notes) + 1)
    ref = next(title_para.iter(q("footnoteReference")))
    old_note = next(n for n in tnotes if n.get(q("id")) == ref.get(q("id")))
    ref.set(q("id"), new_id)
    note = copy.deepcopy(old_note)
    note.set(q("id"), new_id)
    for shd in list(note.iter(q("shd"))):
        shd.getparent().remove(shd)
    for node in note.iter():
        node.attrib.pop(f"{{{W14}}}paraId", None)
    used = {r.get("Id") for r in rels}
    for link in note.iter(q("hyperlink")):
        source = next(r for r in trels if r.get("Id") == link.get(f"{{{R}}}id"))
        rid = next(f"rIdTpl{i}" for i in range(1, 100) if f"rIdTpl{i}" not in used)
        used.add(rid)
        clone = copy.deepcopy(source)
        clone.set("Id", rid)
        rels.append(clone)
        link.set(f"{{{R}}}id", rid)
    notes.append(note)
    for el in front:
        body.remove(el)
    for k, el in enumerate(block):
        body.insert(k, el)
    replacements = {"word/document.xml": etree.tostring(doc, xml_declaration=True, encoding="UTF-8", standalone=True),
                    "word/footnotes.xml": etree.tostring(notes, xml_declaration=True, encoding="UTF-8", standalone=True),
                    "word/_rels/footnotes.xml.rels": etree.tostring(rels, xml_declaration=True, encoding="UTF-8",
                                                                     standalone=True)}
    temp = path.with_suffix(".title.tmp")
    with ZipFile(temp, "w", ZIP_DEFLATED) as z:
        for info, data in entries:
            z.writestr(info, replacements.pop(info.filename, data))
        for name_, data in replacements.items():
            z.writestr(name_, data)
    temp.replace(path)


def validate_final_inputs(results: dict, text: str) -> None:
    """Fail closed on missing author inputs; accept D38 only as the explicitly disclosed deviation it is.

    The original human-validation/account-reconciliation route remains supported. D38 instead requires completed
    AI-review aggregate evidence, an explicit statement that human validation was not completed, and an explicit
    deferred-account declaration with a scoped API-key receipt. It never sets human validation to true.
    """
    def finite(key: str, lo: float, hi: float = float("inf")) -> bool:
        v = results.get(key)
        return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v) and lo <= v <= hi

    def sha(key: str) -> bool:
        return bool(re.fullmatch(r"[0-9a-f]{64}", str(results.get(key, ""))))

    body = " ".join(text.lower().split())
    if "{{" in text or "<!--" in text:
        raise ValueError("Unresolved report placeholders or editorial comments")
    if results.get("meta.preliminary") is not False or results.get("meta.cohort") != "confirmatory" \
            or results.get("meta.analysis_dirty") is not False \
            or not re.fullmatch(r"[0-9a-f]{7,40}", str(results.get("meta.analysis_commit", ""))):
        raise ValueError("Final package requires the confirmatory analysis from clean committed source")
    human = results.get("coding.handcheck.n") == 100 and finite("coding.handcheck.raw_agreement", 0, 1) \
        and finite("coding.handcheck.kappa", -1, 1)
    if human and results.get("review.human.status") == "not performed (D38)":
        raise ValueError("Human agreement fields contradict the declared absence of human validation")
    if not human:
        ai = results.get("review.human.status") == "not performed (D38)" \
            and results.get("review.ai.status") == "completed AI second review (D38)" \
            and results.get("review.ai.reviewer_type") == "AI" \
            and results.get("review.ai.human_validation") is False \
            and results.get("review.ai.blinded_to_classifier_labels") is True \
            and results.get("review.ai.n") == 100 \
            and finite("review.ai.n_agree", 0, 100) \
            and finite("review.ai.raw_agreement", 0, 1) and finite("review.ai.kappa", -1, 1) \
            and sha("review.ai.sheet_sha256") and sha("review.ai.reviewer_metadata_sha256")
        if not ai or not math.isclose(results["review.ai.n_agree"] / 100,
                                      results["review.ai.raw_agreement"], abs_tol=1e-12):
            raise ValueError("Human review is incomplete and no complete, identified D38 AI second review is supplied")
        if not re.search(r"human validation was not (?:completed|performed)", body) \
                or "does not replace human validation" not in body:
            raise ValueError("D38 must explicitly disclose the unmet human-validation requirement")
    account = finite("meta.account_spend_usd", 0) and isinstance(results.get("meta.account_spend_source"), str) \
        and bool(results["meta.account_spend_source"].strip()) \
        and not results["meta.account_spend_source"].strip().lower().startswith("pending")
    if account and results.get("review.account.status") == "deferred (D38)":
        raise ValueError("A whole-account total contradicts the declared deferred reconciliation")
    if not account:
        deferred = results.get("review.account.status") == "deferred (D38)" \
            and finite("costs.key_usage.net_usd", 0) \
            and results.get("costs.key_usage.scope") == "current API key only" \
            and sha("costs.key_usage.source_sha256")
        if not deferred or "reconciliation was deferred" not in body \
                or "no whole-account spending total is claimed" not in body:
            raise ValueError("Account reconciliation is incomplete and no explicit scoped D38 deferral is supplied")
    statement = re.search(r"(?ims)^## LLM Usage Statement\s*\n(.*?)(?=^## |\Z)", text)
    if not statement or len(re.findall(r"[A-Za-z]+", statement.group(1))) < 3:
        raise ValueError("Missing LLM Usage Statement")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--report", type=Path, required=True, help="Filled Markdown report")
    ap.add_argument("--source", type=Path, default=Path("report/REPORT.md"))
    ap.add_argument("--results-dir", type=Path, required=True)
    ap.add_argument("--template", type=Path, required=True)
    ap.add_argument("--out-dir", type=Path, required=True)
    ap.add_argument("--pandoc", default="pandoc")
    ap.add_argument("--tables-url", help="Published directory URL for complete Markdown/CSV tables")
    ap.add_argument("--draft", action="store_true", help="Allow pending fields; filenames identify draft status")
    a = ap.parse_args()
    text = a.report.read_text()
    results = json.loads((a.results_dir / "results.json").read_text())
    if not a.draft:
        validate_final_inputs(results, text)
    a.out_dir.mkdir(parents=True, exist_ok=True)
    printed, records = external_tables(text, a.source.read_text(), a.results_dir / "tables", a.out_dir, a.tables_url)
    # The source uses familiar plain-text exponent notation. Pandoc requires a closing caret or it can consume the
    # fraction's intervening text as one superscript. Delimit only these known price-benchmark superscripts.
    printed = re.sub(r"p̄\^([NM])(?!\^)", r"p̄^\1^", printed)
    start = printed.find("## Appendix A.")
    llm = printed.find("## LLM Usage Statement")
    if start < 0 or llm < start:
        raise SystemExit("Expected Appendix A and final LLM Usage Statement sections")
    main_text = printed[:start].rstrip() + "\n\n" + printed[llm:].strip() + "\n\n" \
        "**Supplement:** Appendices A–H and redacted excerpts are in the supplement. Complete generated tables " \
        "are provided in the accompanying table archive.\n"
    supplement = "# Shared Host Collusion Supplementary Material\n\n" + printed[start:llm].strip() + "\n"
    ending = ".draft" if a.draft else ""
    files = []
    for stem, content in (("MANUSCRIPT", main_text), ("SUPPLEMENT", supplement), ("REPORT", printed)):
        md = a.out_dir / f"{stem}{ending}.md"
        md.write_text(content)
        out = md.with_suffix(".docx")
        subprocess.run([a.pandoc, str(md), "--from", "markdown-implicit_figures-tex_math_dollars-tex_math_single_backslash",
                        "--to", "docx", "--resource-path", f"{a.report.parent.resolve()}:{a.results_dir.resolve()}",
                        "--reference-doc", str(a.template.resolve()), "-o", str(out)], check=True)
        preserve_template_package(out, a.template, a.pandoc)
        format_docx(out, supplement=stem == "SUPPLEMENT", combined=stem == "REPORT")
        if stem != "SUPPLEMENT":
            apply_template_title_block(out, a.template)
        files.append(out.name)
    (a.out_dir / "TABLES_README.md").write_text(
        "# Complete generated tables\n\nThese tables accompany the manuscript and supplement. Each cited generated "
        "table is retained in Markdown and CSV. Long or wide tables are referenced from the supplement rather than "
        "compressed into unreadable print cells. Raw transcripts, human-check data, and attempt-level provider "
        "diagnostics are not included. The Markdown is authoritative for formatting and notes; CSV contains table "
        "cells, including any repeated headers when a source contains multiple tables.\n\n" + "\n".join(
            f"- {r['name']}: {r['rows_including_headers']} rows including headers, up to {r['columns']} columns; {r['placement']}"
                                     for r in records) + "\n")
    with ZipFile(a.out_dir / f"TABLES{ending}.zip", "w", ZIP_DEFLATED) as z:
        z.write(a.out_dir / "TABLES_README.md", "README.md")
        for p in sorted((a.out_dir / "tables").glob("*")):
            z.write(p, "tables/" + p.name)
    manifest = {"status": "draft" if a.draft else "final", "source_report": str(a.report),
                "source_report_sha256": digest(a.report), "template_sha256": digest(a.template),
                "documents": files, "tables": records,
                "analysis_commit": results.get("meta.analysis_commit"),
                "analysis_dirty": results.get("meta.analysis_dirty"),
                "human_review_status": results.get("review.human.status", "human check reported"),
                "ai_review_status": results.get("review.ai.status", "not separately reported"),
                "account_reconciliation_status": results.get("review.account.status", "reconciled amount reported"),
                "published_tables_url": a.tables_url,
                "format": "Apart Letter 1in; Old Standard TT 11pt body; 14pt section headings; visual QA required"}
    (a.out_dir / "package_manifest.json").write_text(json.dumps(manifest, indent=2)+"\n")
    print(json.dumps({"documents": files, "external_tables": sum(r['placement'].startswith('external') for r in records)}, indent=2))


if __name__ == "__main__":
    main()
